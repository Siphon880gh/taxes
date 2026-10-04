#!/usr/bin/env python3
"""Step 2: extract text from every inventoried file.

Backends, first that works wins (everything beyond the standard library is optional):
  PDF     pdfplumber -> PyPDF2/pypdf text layer. If the layer is empty or poor the
          pages are rendered (mutool | pdftoppm | pdf2image) and OCR'd
          (tesseract CLI | pytesseract). Page PNGs are kept in <work>/pages/.
  Images  PIL / sips / magick normalise (HEIC->PNG, upscale small photos, OSD
          rotation fix), then OCR. The PNG is kept for vision review.
  DOCX    python-docx (fallback: unzip word/document.xml). Image-only DOCX ->
          embedded images are exported for vision review.
  DOC/RTF/ODT  textutil (macOS); ODT fallback unzip content.xml
  XLSX    openpyxl. CSV/TSV/TXT/MD/JSON read directly. HTML tags stripped.
  EML     email module: headers + text body; attachments saved to <work>/attachments/.
Files whose text is missing or unreliable are flagged needs_vision so the agent
reads the page images itself.

Usage:
  python3 extract_text.py [INPUT] [--out OUT] [--work WORK] [--only PATH_OR_HASH ...]
                          [--max-pages 6] [--dpi 200] [--force] [--render] [--keep-pii]
"""
import argparse
import email
import email.policy
import html
import re
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

sys.dont_write_bytecode = True  # keep the skill folder free of __pycache__
sys.path.insert(0, str(Path(__file__).resolve().parent))
from taxsort_common import (ensure_work, load_json, mask_pii, now_iso, resolve_paths, run, save_json,
                            tool_available, warn)

TEXT_CAP = 80000  # characters kept per file


# --------------------------------------------------------------------------
# quality
# --------------------------------------------------------------------------

def assess(text: str) -> str:
    """good = enough real words to classify; poor = some text but unreliable (OCR noise); none = nothing usable.

    Tax forms, bills and receipts are short, so the bar is low on word count and
    instead looks at how much of the text is made of pronounceable words.
    """
    words = re.findall(r"[A-Za-z]{2,}", text or "")
    real = [w for w in words if len(w) >= 3 and re.search(r"[aeiouyAEIOUY]", w)]
    nonspace = re.sub(r"\s+", "", text or "")
    if not nonspace or not words:
        return "none"
    alnum = sum(ch.isalnum() for ch in nonspace)
    ratio = alnum / float(len(nonspace))
    real_share = len(real) / float(len(words))
    if len(real) >= 12 and ratio >= 0.55 and real_share >= 0.5:
        return "good"
    if len(real) >= 4 and ratio >= 0.4:
        return "poor"
    return "none"


QUALITY_RANK = {"none": 0, "poor": 1, "good": 2}


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def pdf_text_layer(path: Path, max_pages: int) -> Tuple[List[str], Optional[int], Optional[str]]:
    pages: List[str] = []
    total = None
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            total = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                if max_pages and i >= max_pages:
                    break
                try:
                    pages.append(page.extract_text() or "")
                except Exception:
                    pages.append("")
        return pages, total, "pdfplumber"
    except ImportError:
        pass
    except Exception as exc:  # encrypted / damaged -> try the other library
        warn("pdfplumber failed on %s: %s" % (path.name, exc))
    reader_cls = None
    for mod in ("pypdf", "PyPDF2"):
        try:
            reader_cls = __import__(mod, fromlist=["PdfReader"]).PdfReader
            break
        except Exception:
            continue
    if reader_cls is None:
        return [], total, None
    try:
        reader = reader_cls(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                pass
        total = len(reader.pages)
        pages = []
        for i, page in enumerate(reader.pages):
            if max_pages and i >= max_pages:
                break
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")
        return pages, total, "pypdf"
    except Exception as exc:
        warn("pypdf failed on %s: %s" % (path.name, exc))
        return [], total, None


def pdf_page_count(path: Path) -> Optional[int]:
    if tool_available("mutool"):
        rc, out, err = run(["mutool", "info", str(path)], timeout=60)
        m = re.search(r"Pages:\s*(\d+)", out + err)
        if m:
            return int(m.group(1))
    if tool_available("pdfinfo"):
        rc, out, _ = run(["pdfinfo", str(path)], timeout=60)
        m = re.search(r"Pages:\s*(\d+)", out)
        if m:
            return int(m.group(1))
    return None


def _page_num(p: Path) -> int:
    m = re.search(r"-(\d+)\.png$", p.name)
    return int(m.group(1)) if m else 0


def render_pdf_pages(path: Path, n_pages: int, dpi: int, pages_dir: Path, stem: str) -> List[Path]:
    for old in pages_dir.glob(stem + "-*.png"):
        old.unlink()
    if tool_available("mutool"):
        pattern = pages_dir / (stem + "-%d.png")
        run(["mutool", "draw", "-q", "-o", str(pattern), "-r", str(dpi), str(path), "1-%d" % n_pages], timeout=600)
        files = sorted(pages_dir.glob(stem + "-*.png"), key=_page_num)
        if files:
            return files
    if tool_available("pdftoppm"):
        run(["pdftoppm", "-r", str(dpi), "-png", "-f", "1", "-l", str(n_pages), str(path), str(pages_dir / stem)],
            timeout=600)
        files = sorted(pages_dir.glob(stem + "-*.png"), key=_page_num)
        if files:
            return files
    try:
        from pdf2image import convert_from_path  # type: ignore
        images = convert_from_path(str(path), dpi=dpi, first_page=1, last_page=n_pages)
        files = []
        for i, img in enumerate(images, 1):
            target = pages_dir / ("%s-%d.png" % (stem, i))
            img.save(str(target))
            files.append(target)
        return files
    except Exception:
        return []


# --------------------------------------------------------------------------
# OCR
# --------------------------------------------------------------------------

def ocr_available() -> bool:
    if tool_available("tesseract"):
        return True
    try:
        import pytesseract  # type: ignore  # noqa: F401
        return True
    except Exception:
        return False


def renderer_available() -> bool:
    if tool_available("mutool") or tool_available("pdftoppm"):
        return True
    try:
        import pdf2image  # type: ignore  # noqa: F401
        return True
    except Exception:
        return False


def detect_rotation(png: Path) -> int:
    if not tool_available("tesseract"):
        return 0
    rc, out, err = run(["tesseract", str(png), "stdout", "--psm", "0"], timeout=90)
    m = re.search(r"Rotate:\s*(\d+)", out + err)
    return int(m.group(1)) % 360 if m else 0


def rotate_png(png: Path, degrees: int) -> Path:
    try:
        from PIL import Image  # type: ignore
        img = Image.open(str(png))
        # tesseract reports the clockwise rotation needed; PIL rotates counter-clockwise.
        img = img.rotate(-degrees, expand=True)
        img.save(str(png))
    except Exception as exc:
        warn("could not rotate %s: %s" % (png.name, exc))
    return png


def ocr_png(png: Path) -> Optional[str]:
    if tool_available("tesseract"):
        rc, out, err = run(["tesseract", str(png), "stdout", "-l", "eng", "--psm", "3"], timeout=300)
        if rc == 0:
            return out
        warn("tesseract failed on %s: %s" % (png.name, err.strip()[:200]))
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
        return pytesseract.image_to_string(Image.open(str(png)))
    except Exception:
        return None


def ocr_pages(pngs: List[Path], fix_rotation: bool = True) -> List[str]:
    texts = []
    for png in pngs:
        if fix_rotation:
            rot = detect_rotation(png)
            if rot in (90, 180, 270):
                rotate_png(png, rot)
        texts.append(ocr_png(png) or "")
    return texts


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------

def prepare_image(path: Path, pages_dir: Path, stem: str) -> Optional[Path]:
    """Produce a PNG copy suitable for OCR and vision; returns None if nothing could read the image."""
    target = pages_dir / (stem + "-1.png")
    ext = path.suffix.lower()
    src_for_pil = path
    if ext in (".heic", ".heif"):
        converted = None
        try:
            import pillow_heif  # type: ignore
            pillow_heif.register_heif_opener()
            from PIL import Image  # type: ignore
            Image.open(str(path)).convert("RGB").save(str(target))
            converted = target
        except Exception:
            pass
        if converted is None and tool_available("sips"):
            rc, _, _ = run(["sips", "-s", "format", "png", str(path), "--out", str(target)], timeout=120)
            if rc == 0 and target.exists():
                converted = target
        if converted is None and tool_available("magick"):
            rc, _, _ = run(["magick", str(path), str(target)], timeout=120)
            if rc == 0 and target.exists():
                converted = target
        if converted is None:
            return None
        src_for_pil = converted
    try:
        from PIL import Image, ImageOps  # type: ignore
        img = Image.open(str(src_for_pil))
        img = ImageOps.exif_transpose(img)
        img = img.convert("RGB")
        longest = max(img.size)
        if longest < 1600:
            scale = min(2.5, 1600.0 / max(longest, 1))
            img = img.resize((int(img.width * scale), int(img.height * scale)))
        img.save(str(target))
        return target
    except Exception:
        if src_for_pil != path:
            return src_for_pil
        if tool_available("magick"):
            rc, _, _ = run(["magick", str(path), str(target)], timeout=120)
            if rc == 0 and target.exists():
                return target
        if tool_available("sips"):
            rc, _, _ = run(["sips", "-s", "format", "png", str(path), "--out", str(target)], timeout=120)
            if rc == 0 and target.exists():
                return target
        if ext in (".png", ".jpg", ".jpeg"):
            return path
        return None


# --------------------------------------------------------------------------
# Office / text / email
# --------------------------------------------------------------------------

def strip_tags(markup: str) -> str:
    markup = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", markup)
    markup = re.sub(r"(?i)<\s*(br|/p|/div|/tr|/li|/h\d)\s*/?>", "\n", markup)
    markup = re.sub(r"(?i)<\s*/t[dh]\s*>", "\t", markup)
    text = re.sub(r"<[^>]+>", " ", markup)
    text = html.unescape(text)
    return re.sub(r"[ \t]+", " ", text)


def docx_text(path: Path, pages_dir: Path, stem: str) -> Tuple[str, List[Path]]:
    text = ""
    try:
        import docx  # type: ignore
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(c.text for c in row.cells))
        for section in d.sections:
            try:
                parts.extend(p.text for p in section.header.paragraphs)
                parts.extend(p.text for p in section.footer.paragraphs)
            except Exception:
                pass
        text = "\n".join(parts)
    except Exception:
        try:
            with zipfile.ZipFile(str(path)) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", "ignore")
            xml = re.sub(r"</w:p>", "\n", xml)
            xml = re.sub(r"<w:tab/>", "\t", xml)
            text = re.sub(r"<[^>]+>", "", xml)
        except Exception:
            text = ""
    images: List[Path] = []
    if assess(text) != "good":
        try:
            with zipfile.ZipFile(str(path)) as zf:
                media = [n for n in zf.namelist() if n.startswith("word/media/")
                         and n.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"))]
                for i, name in enumerate(media[:4], 1):
                    target = pages_dir / ("%s-%d%s" % (stem, i, Path(name).suffix.lower()))
                    with zf.open(name) as src, open(target, "wb") as dst:
                        shutil.copyfileobj(src, dst)
                    images.append(target)
        except Exception:
            pass
    return text, images


def textutil_text(path: Path) -> str:
    if tool_available("textutil"):
        rc, out, _ = run(["textutil", "-convert", "txt", "-stdout", str(path)], timeout=120)
        if rc == 0 and out.strip():
            return out
    if path.suffix.lower() == ".odt":
        try:
            with zipfile.ZipFile(str(path)) as zf:
                xml = zf.read("content.xml").decode("utf-8", "ignore")
            xml = re.sub(r"</text:p>", "\n", xml)
            return re.sub(r"<[^>]+>", "", xml)
        except Exception:
            return ""
    if path.suffix.lower() == ".rtf":
        try:
            raw = path.read_text(encoding="latin-1", errors="ignore")
            raw = re.sub(r"\\par[d]?", "\n", raw)
            raw = re.sub(r"\\'[0-9a-f]{2}", "", raw)
            raw = re.sub(r"\\[a-z]+-?\d* ?", "", raw)
            return re.sub(r"[{}]", "", raw)
        except Exception:
            return ""
    return ""


def sheet_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # type: ignore
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append("## Sheet: %s" % ws.title)
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 400:
                        parts.append("... (truncated)")
                        break
                    cells = ["" if v is None else str(v) for v in row]
                    if any(c.strip() for c in cells):
                        parts.append("\t".join(cells).rstrip())
            return "\n".join(parts)
        except Exception as exc:
            warn("openpyxl failed on %s: %s" % (path.name, exc))
            return ""
    if ext == ".xls":
        try:
            import xlrd  # type: ignore
            book = xlrd.open_workbook(str(path))
            parts = []
            for sh in book.sheets():
                parts.append("## Sheet: %s" % sh.name)
                for r in range(min(sh.nrows, 400)):
                    parts.append("\t".join(str(c) for c in sh.row_values(r)))
            return "\n".join(parts)
        except Exception:
            return ""
    if ext == ".ods":
        try:
            with zipfile.ZipFile(str(path)) as zf:
                xml = zf.read("content.xml").decode("utf-8", "ignore")
            xml = re.sub(r"</table:table-row>", "\n", xml)
            xml = re.sub(r"</table:table-cell>", "\t", xml)
            return re.sub(r"<[^>]+>", "", xml)
        except Exception:
            return ""
    return ""


def plain_text(path: Path) -> str:
    raw = path.read_bytes()[: TEXT_CAP * 2]
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("latin-1", "ignore")
    if path.suffix.lower() in (".html", ".htm"):
        text = strip_tags(text)
    return text


def eml_text(path: Path, attach_dir: Path) -> Tuple[str, List[str]]:
    with open(str(path), "rb") as fh:
        msg = email.message_from_binary_file(fh, policy=email.policy.default)
    parts = []
    for hdr in ("From", "To", "Subject", "Date"):
        if msg.get(hdr):
            parts.append("%s: %s" % (hdr, msg.get(hdr)))
    attachments: List[str] = []
    body_plain, body_html = "", ""
    for part in msg.walk():
        ctype = part.get_content_type()
        fname = part.get_filename()
        if fname:
            attach_dir.mkdir(parents=True, exist_ok=True)
            target = attach_dir / Path(fname).name
            try:
                payload = part.get_payload(decode=True)
                if payload:
                    target.write_bytes(payload)
                    attachments.append(str(target))
            except Exception:
                pass
            continue
        if ctype == "text/plain" and not body_plain:
            try:
                body_plain = part.get_content()
            except Exception:
                pass
        elif ctype == "text/html" and not body_html:
            try:
                body_html = strip_tags(part.get_content())
            except Exception:
                pass
    parts.append("")
    parts.append(body_plain or body_html)
    if attachments:
        parts.append("")
        parts.append("Attachments: " + ", ".join(Path(a).name for a in attachments))
    return "\n".join(parts), attachments


def zip_listing(path: Path) -> str:
    try:
        with zipfile.ZipFile(str(path)) as zf:
            names = zf.namelist()
        return "ZIP archive with %d entries:\n%s" % (len(names), "\n".join(names[:60]))
    except Exception:
        return ""


# --------------------------------------------------------------------------
# driver
# --------------------------------------------------------------------------

def extract_one(f: Dict, work: Path, max_pages: int, dpi: int, force_render: bool) -> Dict:
    path = Path(f["path"])
    stem = f["hash"][:16]
    pages_dir = work / "pages"
    rec: Dict = {
        "path": f["path"], "rel": f["rel"], "hash": f["hash"], "kind": f["kind"], "method": "none",
        "quality": "none", "chars": 0, "words": 0, "pages_total": None, "pages_sampled": 0,
        "page_images": [], "attachments": [], "needs_vision": False, "vision_reason": "", "error": None,
        "extracted_at": now_iso(),
    }
    text = ""
    kind, ext = f["kind"], f["ext"]
    try:
        if kind == "pdf":
            pages, total, backend = pdf_text_layer(path, max_pages)
            rec["pages_total"] = total if total is not None else pdf_page_count(path)
            rec["pages_sampled"] = len(pages)
            layer = "\n\n".join("--- page %d ---\n%s" % (i + 1, t) for i, t in enumerate(pages) if t.strip())
            layer_q = assess(layer)
            text, rec["method"] = layer, ("pdf-text" if backend else "none")
            if layer_q != "good" or force_render:
                n = rec["pages_total"] or max_pages or 6
                if max_pages:
                    n = min(n, max_pages)
                pngs = render_pdf_pages(path, max(1, n), dpi, pages_dir, stem) if renderer_available() else []
                rec["page_images"] = [str(p) for p in pngs]
                if layer_q != "good":
                    if not pngs:
                        rec["needs_vision"] = True
                        rec["vision_reason"] = ("text layer %s and no PDF renderer available "
                                                "(brew install mupdf-tools or poppler)" % layer_q)
                    elif not ocr_available():
                        rec["needs_vision"] = True
                        rec["vision_reason"] = "text layer %s and no OCR engine (brew install tesseract)" % layer_q
                    else:
                        ocr_texts = ocr_pages(pngs)
                        ocr = "\n\n".join("--- page %d ---\n%s" % (i + 1, t) for i, t in enumerate(ocr_texts) if t.strip())
                        if QUALITY_RANK[assess(ocr)] >= QUALITY_RANK[layer_q]:
                            text, rec["method"] = ocr, "pdf-ocr"
                            rec["pages_sampled"] = len(pngs)
                        if assess(text) != "good":
                            rec["needs_vision"] = True
                            rec["vision_reason"] = "OCR result is %s; read the page images" % assess(text)
        elif kind == "image":
            png = prepare_image(path, pages_dir, stem)
            if png is None:
                rec["needs_vision"] = True
                rec["vision_reason"] = "image format could not be converted (HEIC needs sips/magick/pillow-heif)"
                rec["page_images"] = [str(path)]
            else:
                rec["page_images"] = [str(png)]
                if ocr_available():
                    text = "\n".join(ocr_pages([png]))
                    rec["method"] = "image-ocr"
                    if assess(text) != "good":
                        rec["needs_vision"] = True
                        rec["vision_reason"] = "OCR result is %s; read the image" % assess(text)
                else:
                    rec["needs_vision"] = True
                    rec["vision_reason"] = "no OCR engine (brew install tesseract); read the image"
        elif ext == ".docx":
            text, images = docx_text(path, pages_dir, stem)
            rec["method"] = "docx"
            rec["page_images"] = [str(p) for p in images]
            if assess(text) != "good":
                if images and ocr_available():
                    text = text + "\n" + "\n".join(ocr_pages(images))
                    rec["method"] = "docx+ocr"
                if assess(text) != "good":
                    rec["needs_vision"] = True
                    rec["vision_reason"] = "little text in document%s" % ("; embedded images exported" if images else "")
        elif ext in (".doc", ".rtf", ".odt", ".pages"):
            text = textutil_text(path)
            rec["method"] = "textutil" if text else "unsupported"
            if assess(text) == "none":
                rec["needs_vision"] = True
                rec["vision_reason"] = "no text extracted (export to PDF or DOCX, or read manually)"
        elif ext in (".xlsx", ".xlsm", ".xls", ".ods"):
            text = sheet_text(path)
            rec["method"] = "sheet" if text else "unsupported"
            if not text:
                rec["needs_vision"] = True
                rec["vision_reason"] = "spreadsheet could not be read (install openpyxl / xlrd, or export CSV)"
        elif ext == ".numbers":
            rec["method"] = "unsupported"
            rec["needs_vision"] = True
            rec["vision_reason"] = "Apple Numbers file: export to XLSX or CSV"
        elif kind == "text" or ext in (".csv", ".tsv"):
            text = plain_text(path)
            rec["method"] = "text"
        elif ext == ".eml":
            text, attachments = eml_text(path, work / "attachments" / stem)
            rec["method"] = "eml"
            rec["attachments"] = attachments
        elif ext == ".msg":
            try:
                import extract_msg  # type: ignore
                m = extract_msg.Message(str(path))
                text = "From: %s\nSubject: %s\nDate: %s\n\n%s" % (m.sender, m.subject, m.date, m.body)
                rec["method"] = "msg"
            except Exception:
                rec["method"] = "unsupported"
                rec["needs_vision"] = True
                rec["vision_reason"] = "Outlook .msg needs the extract_msg package (pip install extract-msg)"
        elif kind == "archive":
            text = zip_listing(path) if ext == ".zip" else ""
            rec["method"] = "archive"
            rec["vision_reason"] = "archive: extract its contents into the inbox"
        else:
            rec["method"] = "unsupported"
            rec["needs_vision"] = True
            rec["vision_reason"] = "unsupported file type %s" % ext
    except Exception as exc:  # never let one file stop the batch
        rec["error"] = "%s: %s" % (type(exc).__name__, exc)
        rec["needs_vision"] = True
        rec["vision_reason"] = "extraction error; read the file manually"

    text = (text or "")[:TEXT_CAP]
    rec["quality"] = assess(text)
    rec["chars"] = len(text)
    rec["words"] = len(re.findall(r"[A-Za-z]{2,}", text))
    rec["_text"] = text
    return rec


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", nargs="?")
    ap.add_argument("--out")
    ap.add_argument("--work")
    ap.add_argument("--only", nargs="*", default=[], help="limit to files whose path contains / hash starts with these")
    ap.add_argument("--max-pages", type=int, default=6, help="pages to read per PDF (0 = all)")
    ap.add_argument("--dpi", type=int, default=200, help="render resolution for OCR / vision")
    ap.add_argument("--force", action="store_true", help="re-extract even if cached")
    ap.add_argument("--render", action="store_true", help="also render page images for PDFs that have good text")
    ap.add_argument("--keep-pii", action="store_true", help="do not mask SSNs in the cached text")
    args = ap.parse_args()

    paths = resolve_paths(args.input, args.out, args.work)
    work = ensure_work(paths["work"])
    inventory = load_json(work / "inventory.json")
    if not inventory:
        rc, out, err = run([sys.executable, str(Path(__file__).resolve().parent / "inventory.py"),
                            str(paths["input"]), "--out", str(paths["out"]), "--work", str(work)])
        inventory = load_json(work / "inventory.json")
        if not inventory:
            sys.stderr.write(err)
            return 1
    extracted: Dict[str, Dict] = load_json(work / "extracted.json", {}) or {}

    files = inventory["files"]
    if args.only:
        sel = []
        for f in files:
            for token in args.only:
                if token in f["path"] or token in f["rel"] or f["hash"].startswith(token):
                    sel.append(f)
                    break
        files = sel
    todo = []
    for f in files:
        prev = extracted.get(f["hash"])
        if prev and not args.force and not prev.get("error") and prev.get("path") == f["path"]:
            if not args.render or prev.get("page_images"):
                continue
        todo.append(f)
    print("Extracting %d of %d files (cached: %d)" % (len(todo), len(files), len(files) - len(todo)))

    text_dir = work / "text"
    for i, f in enumerate(todo, 1):
        sys.stdout.write("  [%d/%d] %s ... " % (i, len(todo), f["rel"]))
        sys.stdout.flush()
        if f.get("duplicate_of"):
            primary = next((x for x in inventory["files"] if x["rel"] == f["duplicate_of"]), None)
            if primary and primary["hash"] in extracted:
                print("duplicate, reusing text")
                continue
        rec = extract_one(f, work, args.max_pages, args.dpi, args.render)
        text = rec.pop("_text")
        if not args.keep_pii:
            text = mask_pii(text)
        (text_dir / (f["hash"] + ".txt")).write_text(text, encoding="utf-8")
        rec["text_file"] = str(text_dir / (f["hash"] + ".txt"))
        extracted[f["hash"]] = rec
        save_json(work / "extracted.json", extracted)
        flag = " NEEDS VISION (%s)" % rec["vision_reason"] if rec["needs_vision"] else ""
        print("%s, %s, %d words%s" % (rec["method"], rec["quality"], rec["words"], flag))

    # summary over everything in inventory
    recs = [extracted.get(f["hash"]) for f in inventory["files"]]
    recs = [r for r in recs if r]
    from collections import Counter
    print("")
    print("Methods : %s" % ", ".join("%s %d" % kv for kv in sorted(Counter(r["method"] for r in recs).items())))
    print("Quality : %s" % ", ".join("%s %d" % kv for kv in sorted(Counter(r["quality"] for r in recs).items())))
    vision = [r for r in recs if r["needs_vision"]]
    if vision:
        print("")
        print("NEEDS VISION (%d) - open these images with the Read tool and classify them yourself:" % len(vision))
        for r in vision:
            print("  - %s: %s" % (r["rel"], r["vision_reason"]))
            for img in r["page_images"][:6]:
                print("      %s" % img)
    attach = [r for r in recs if r.get("attachments")]
    if attach:
        print("")
        print("Email attachments were saved under %s; copy the relevant ones into the inbox to sort them." % (work / "attachments"))
    errors = [r for r in recs if r.get("error")]
    for r in errors:
        print("ERROR %s: %s" % (r["rel"], r["error"]))
    print("")
    print("Next: python3 %s/classify.py \"%s\" --out \"%s\"" % (Path(__file__).resolve().parent, paths["input"], paths["out"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
